"""Importar datos desde un archivo de Excel sin formato predefinido."""

import io

import pytest
from openpyxl import Workbook

from tests.conftest import API, cabeceras, unico
from tests.test_multicuenta import crear_cuenta_completa


def excel(encabezados: list[str], filas: list[list]) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.append(["Inventario de la finca"])  # fila suelta antes del encabezado
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


@pytest.fixture(scope="module")
def finca(cliente, token_plataforma):
    datos = crear_cuenta_completa(cliente, token_plataforma, "CuentaImport")
    cab = cabeceras(datos["token"], finca_id=datos["finca_id"])

    bodega = cliente.post(f"{API}/bodegas", headers=cab, json={"codigo": "B1", "nombre": "Bodega"}).json()
    galpon = cliente.post(
        f"{API}/galpones", headers=cab, json={"codigo": "G1", "nombre": "Galpon 1", "capacidad": 1000}
    ).json()
    lote = cliente.post(
        f"{API}/lotes",
        headers=cab,
        json={"codigo": "L1", "galpon_id": galpon["id"], "fecha_ingreso": "2026-01-05", "aves_iniciales": 500},
    ).json()

    datos.update({"cab": cab, "bodega": bodega, "galpon": galpon, "lote": lote})
    return datos


def subir(cliente, finca, tipo: str, contenido: bytes, nombre="archivo.xlsx"):
    return cliente.post(
        f"{API}/importacion/analizar",
        headers=finca["cab"],
        data={"tipo": tipo},
        files={"archivo": (nombre, contenido, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_los_tipos_traen_sus_campos(cliente, finca):
    tipos = cliente.get(f"{API}/importacion/tipos", headers=finca["cab"]).json()
    claves = {t["clave"] for t in tipos}
    assert {"articulos", "entrada_inventario", "proveedores", "sanidad", "produccion"} <= claves
    articulos = next(t for t in tipos if t["clave"] == "articulos")
    assert any(c["clave"] == "codigo" and c["obligatorio"] for c in articulos["campos"])


def test_adivina_las_columnas_aunque_tengan_otro_nombre(cliente, finca):
    contenido = excel(
        ["CÓDIGO", "DESCRIPCIÓN DEL PRODUCTO", "UND", "CANTIDAD", "VR UNITARIO", "PROVEEDOR", "FECHA"],
        [["AL-01", "Concentrado ponedora", "kg", 500, 2500, "Agroinsumos SAS", "2026-02-01"]],
    )
    respuesta = subir(cliente, finca, "entrada_inventario", contenido)
    assert respuesta.status_code == 201, respuesta.text
    analisis = respuesta.json()

    assert analisis["filas_totales"] == 1
    assert "CÓDIGO" in analisis["columnas"]
    sugerido = analisis["mapeo_sugerido"]
    assert sugerido["codigo"] == "CÓDIGO"
    assert sugerido["articulo"] == "DESCRIPCIÓN DEL PRODUCTO"
    assert sugerido["cantidad"] == "CANTIDAD"
    assert sugerido["costo_unitario"] == "VR UNITARIO"
    assert sugerido["proveedor"] == "PROVEEDOR"
    finca["analisis_entrada"] = analisis


def test_validar_muestra_los_errores_sin_guardar_nada(cliente, finca):
    contenido = excel(
        ["Codigo", "Articulo", "Cantidad", "Costo"],
        [
            ["AL-01", "Concentrado ponedora", 500, 2500],
            ["VI-01", "Vacuna Newcastle", "mil", 900],   # cantidad invalida
            ["", "", "", ""],                              # fila vacia se ignora al leer
            ["AL-02", "", 100, 1000],                      # falta el articulo
        ],
    )
    analisis = subir(cliente, finca, "entrada_inventario", contenido).json()
    finca["importacion"] = analisis["id"]

    validacion = cliente.post(
        f"{API}/importacion/{analisis['id']}/validar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {"bodega_id": finca["bodega"]["id"]}},
    )
    assert validacion.status_code == 200, validacion.text
    datos = validacion.json()
    assert datos["filas_ok"] == 1
    assert datos["filas_error"] == 2
    assert any("numero" in error for error in datos["errores"])

    # todavia no hay nada en el inventario
    assert cliente.get(f"{API}/articulos", headers=finca["cab"]).json() == []


def test_falta_indicar_una_columna_obligatoria(cliente, finca):
    respuesta = cliente.post(
        f"{API}/importacion/{finca['importacion']}/validar",
        headers=finca["cab"],
        json={"mapeo": {"cantidad": "Cantidad"}, "opciones": {}},
    )
    assert respuesta.status_code == 400
    assert "Articulo" in respuesta.json()["detail"]["mensaje"]


def test_aplicar_crea_articulos_proveedor_y_entrada(cliente, finca):
    contenido = excel(
        ["Codigo", "Articulo", "Unidad", "Cantidad", "Costo", "Proveedor", "Fecha", "Factura"],
        [
            ["AL-01", "Concentrado ponedora", "kg", 500, 2500, "Agroinsumos SAS", "2026-02-01", "FAC-100"],
            ["VI-01", "Vacuna Newcastle", "dosis", 1000, 900, "Agroinsumos SAS", "2026-02-01", "FAC-100"],
        ],
    )
    analisis = subir(cliente, finca, "entrada_inventario", contenido).json()
    cuerpo = {
        "mapeo": analisis["mapeo_sugerido"],
        "opciones": {"bodega_id": finca["bodega"]["id"], "crear_articulos": True},
    }

    aplicado = cliente.post(f"{API}/importacion/{analisis['id']}/aplicar", headers=finca["cab"], json=cuerpo)
    assert aplicado.status_code == 200, aplicado.text
    datos = aplicado.json()
    assert datos["estado"] == "aplicada"
    assert datos["filas_ok"] == 2
    assert datos["resumen"]["movimientos"] == 1  # misma fecha y factura: un solo movimiento

    articulos = cliente.get(f"{API}/articulos", headers=finca["cab"]).json()
    assert {a["codigo"] for a in articulos} == {"AL-01", "VI-01"}
    assert next(a for a in articulos if a["codigo"] == "AL-01")["existencia_total"] == 500

    proveedores = cliente.get(f"{API}/proveedores", headers=finca["cab"]).json()
    assert any(p["nombre"] == "Agroinsumos SAS" for p in proveedores)

    finca["importacion_aplicada"] = analisis["id"]


def test_no_se_aplica_dos_veces(cliente, finca):
    respuesta = cliente.post(
        f"{API}/importacion/{finca['importacion_aplicada']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": {}, "opciones": {}},
    )
    assert respuesta.status_code == 400


def test_deshacer_devuelve_el_inventario(cliente, finca):
    respuesta = cliente.post(
        f"{API}/importacion/{finca['importacion_aplicada']}/revertir", headers=finca["cab"], json={}
    )
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.json()["estado"] == "revertida"

    # los articulos que se crearon con la importacion quedan desactivados
    articulos = cliente.get(f"{API}/articulos", headers=finca["cab"]).json()
    assert articulos == []
    inactivos = cliente.get(f"{API}/articulos?incluir_inactivos=true", headers=finca["cab"]).json()
    assert all(a["existencia_total"] == 0 for a in inactivos)

    movimientos = cliente.get(f"{API}/movimientos", headers=finca["cab"]).json()
    assert all(m["anulado"] for m in movimientos["datos"])


def test_importar_produccion_de_huevos(cliente, finca):
    contenido = excel(
        ["Dia", "Lote", "Tamaño", "Huevos"],
        [
            ["2026-03-01", "L1", "AAA", 320],
            ["2026-03-01", "L1", "AA", 150],
            ["2026-03-02", "L9", "AAA", 100],  # lote que no existe
        ],
    )
    analisis = subir(cliente, finca, "produccion", contenido).json()
    assert analisis["mapeo_sugerido"]["tipo_huevo"] == "Tamaño"

    aplicado = cliente.post(
        f"{API}/importacion/{analisis['id']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {}},
    ).json()
    assert aplicado["filas_ok"] == 2
    assert aplicado["filas_error"] == 1

    stock = {f["tipo"]: f["cantidad"] for f in cliente.get(f"{API}/stock-huevos", headers=finca["cab"]).json()}
    assert stock["AAA"] == 320
    assert stock["AA"] == 150

    # al deshacer, el stock vuelve a cero
    cliente.post(f"{API}/importacion/{analisis['id']}/revertir", headers=finca["cab"], json={})
    stock = {f["tipo"]: f["cantidad"] for f in cliente.get(f"{API}/stock-huevos", headers=finca["cab"]).json()}
    assert stock.get("AAA", 0) == 0


def test_importar_vacunas(cliente, finca):
    contenido = excel(
        ["Fecha", "Vacuna", "Lote", "Via", "Dosis", "Aves", "Aplicado por"],
        [["2026-03-05", "Newcastle La Sota", "L1", "ocular", "1 por ave", 480, "Juan"]],
    )
    analisis = subir(cliente, finca, "sanidad", contenido).json()
    aplicado = cliente.post(
        f"{API}/importacion/{analisis['id']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {}},
    ).json()
    assert aplicado["filas_ok"] == 1

    registros = cliente.get(f"{API}/sanidad", headers=finca["cab"]).json()
    assert any(r["producto"] == "Newcastle La Sota" and r["via"] == "ocular" for r in registros)


def test_leer_un_csv_con_punto_y_coma(cliente, finca):
    csv = "nombre;nit;telefono\nDistribuidora El Campo;900123;3101234567\n".encode("utf-8")
    respuesta = cliente.post(
        f"{API}/importacion/analizar",
        headers=finca["cab"],
        data={"tipo": "proveedores"},
        files={"archivo": ("proveedores.csv", csv, "text/csv")},
    )
    assert respuesta.status_code == 201, respuesta.text
    analisis = respuesta.json()
    assert analisis["columnas"] == ["nombre", "nit", "telefono"]
    assert analisis["mapeo_sugerido"]["documento"] == "nit"

    aplicado = cliente.post(
        f"{API}/importacion/{analisis['id']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {}},
    ).json()
    assert aplicado["filas_ok"] == 1


def test_guardar_y_usar_una_plantilla(cliente, finca):
    plantilla = cliente.post(
        f"{API}/importacion/plantillas",
        headers=finca["cab"],
        json={
            "tipo": "entrada_inventario",
            "nombre": unico("Formato")[:80],
            "mapeo": {"articulo": "Articulo", "cantidad": "Cantidad"},
        },
    )
    assert plantilla.status_code == 201, plantilla.text

    guardadas = cliente.get(f"{API}/importacion/plantillas?tipo=entrada_inventario", headers=finca["cab"]).json()
    assert any(p["id"] == plantilla.json()["id"] for p in guardadas)

    # las plantillas aparecen al subir un archivo del mismo tipo
    analisis = subir(cliente, finca, "entrada_inventario", excel(["Articulo", "Cantidad"], [["X", 1]])).json()
    assert any(p["id"] == plantilla.json()["id"] for p in analisis["plantillas"])


def test_no_se_lee_un_formato_raro(cliente, finca):
    respuesta = cliente.post(
        f"{API}/importacion/analizar",
        headers=finca["cab"],
        data={"tipo": "articulos"},
        files={"archivo": ("datos.pdf", b"%PDF-1.4 algo", "application/pdf")},
    )
    assert respuesta.status_code == 400
    assert "xlsx" in respuesta.json()["detail"]["mensaje"]


def test_otra_cuenta_no_ve_las_importaciones(cliente, finca, token_plataforma):
    otra = crear_cuenta_completa(cliente, token_plataforma, "CuentaImportB")
    cab = cabeceras(otra["token"], finca_id=otra["finca_id"])

    assert cliente.get(f"{API}/importacion", headers=cab).json() == []
    assert cliente.get(f"{API}/importacion/{finca['importacion_aplicada']}", headers=cab).status_code == 404


def test_una_finca_no_ve_ni_aplica_importaciones_de_otra(cliente, finca):
    archivo = excel(["Codigo", "Nombre", "Unidad"], [[unico("X")[:20], "Articulo de la finca 1", "kg"]])
    subida = subir(cliente, finca, "articulos", archivo)
    assert subida.status_code == 201, subida.text
    importacion_id = subida.json()["id"]

    admin = cabeceras(finca["token"])
    finca2 = cliente.post(f"{API}/fincas", headers=admin, json={"codigo": "P2", "nombre": unico("Otra")}).json()
    email = f"{unico('sup2')}@avicola.com"
    cliente.post(
        f"{API}/usuarios",
        headers=admin,
        json={
            "nombres": "Supervisor",
            "email": email,
            "rol": "supervisor",
            "clave": "ClaveDePrueba123",
            "debe_cambiar_clave": False,
            "fincas": [{"finca_id": finca2["id"], "solo_lectura": False}],
        },
    )
    token = cliente.post(f"{API}/auth/login", json={"email": email, "clave": "ClaveDePrueba123"}).json()["token"]
    cab = cabeceras(token, finca2["id"])

    assert cliente.get(f"{API}/importacion/{importacion_id}", headers=cab).status_code == 404
    assert cliente.post(f"{API}/importacion/{importacion_id}/validar", headers=cab, json={}).status_code in (404, 422)
    historial = cliente.get(f"{API}/importacion", headers=cab).json()
    assert importacion_id not in {i["id"] for i in historial}


def test_un_archivo_de_mas_de_5_mb_se_rechaza(cliente, finca):
    grande = b"codigo;nombre\n" + b"A;" + b"x" * (5 * 1024 * 1024 + 10) + b"\n"
    respuesta = subir(cliente, finca, "articulos", grande, nombre="grande.csv")
    assert respuesta.status_code == 400
    assert "5 MB" in respuesta.json()["detail"]["mensaje"]


def test_un_zip_que_no_es_excel_se_rechaza(cliente, finca):
    respuesta = subir(cliente, finca, "articulos", b"esto no es un excel", nombre="falso.xlsx")
    assert respuesta.status_code == 400
    assert "xlsx" in respuesta.json()["detail"]["mensaje"]


def test_una_bomba_de_descompresion_se_rechaza(cliente, finca):
    import zipfile

    memoria = io.BytesIO()
    with zipfile.ZipFile(memoria, "w", zipfile.ZIP_DEFLATED) as bomba:
        bomba.writestr("xl/worksheets/sheet1.xml", b"0" * (80 * 1024 * 1024))
    respuesta = subir(cliente, finca, "articulos", memoria.getvalue(), nombre="bomba.xlsx")
    assert respuesta.status_code == 400
    assert "demasiado grande" in respuesta.json()["detail"]["mensaje"]


def test_una_hoja_con_un_millon_de_filas_vacias_no_se_recorre_entera(cliente, finca):
    import time

    from app.servicios.importacion import MAX_FILAS_RECORRIDAS

    libro = Workbook()
    hoja = libro.active
    hoja.append(["Codigo", "Nombre"])
    hoja.append(["A1", "Primero"])
    hoja.cell(row=MAX_FILAS_RECORRIDAS + 500, column=1, value="Escondido")
    hoja.cell(row=MAX_FILAS_RECORRIDAS + 500, column=2, value="Muy abajo")
    memoria = io.BytesIO()
    libro.save(memoria)

    inicio = time.monotonic()
    respuesta = subir(cliente, finca, "articulos", memoria.getvalue())
    assert time.monotonic() - inicio < 20
    assert respuesta.status_code == 201, respuesta.text
    assert respuesta.json()["filas_totales"] == 1


@pytest.mark.parametrize(
    "valor, esperado",
    [
        ("1.500", "1500"),
        ("1.234,56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("2,5", "2.5"),
        ("2.5", "2.5"),
        ("1,500,000", "1500000"),
        ("$ 18.000", "18000"),
        (1.125, "1.125"),
        (40, "40"),
    ],
)
def test_lee_los_numeros_como_se_escriben_en_colombia(valor, esperado):
    from decimal import Decimal

    from app.servicios.importacion import numero

    assert numero(valor, "Cantidad") == Decimal(esperado)


@pytest.mark.parametrize("valor", ["NaN", "Infinity", "-inf", "1e999999", "abc", True])
def test_rechaza_lo_que_no_es_un_numero(valor):
    from app.servicios.importacion import numero

    with pytest.raises(ValueError):
        numero(valor, "Cantidad")


def test_un_csv_de_excel_en_windows_conserva_las_tildes(cliente, finca):
    contenido = "Codigo;Nombre;Unidad\nV1;Vacuna Gumboro dosis única;dosis\n".encode("cp1252")
    respuesta = subir(cliente, finca, "articulos", contenido, nombre="vacunas.csv")
    assert respuesta.status_code == 201, respuesta.text
    assert respuesta.json()["vista_previa"][0]["Nombre"] == "Vacuna Gumboro dosis única"


def test_una_bodega_que_no_es_numero_no_rompe_la_importacion(cliente, finca):
    contenido = excel(["Articulo", "Cantidad"], [["Concentrado", 10]])
    analisis = subir(cliente, finca, "entrada_inventario", contenido).json()
    respuesta = cliente.post(
        f"{API}/importacion/{analisis['id']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {"bodega_id": "la central"}},
    )
    assert respuesta.status_code == 400


def test_revertir_deja_el_dia_como_estaba_antes_de_importar(cliente, finca):
    from sqlalchemy import select

    from app.core.db import SesionLocal
    from app.modelos.aves import ProduccionHuevos

    def stock_aaa():
        filas = cliente.get(f"{API}/stock-huevos", headers=finca["cab"]).json()
        return next((f["cantidad"] for f in filas if f["tipo"] == "AAA"), 0)

    aaa = next(t for t in cliente.get(f"{API}/tipos-huevo", headers=finca["cab"]).json() if t["nombre"] == "AAA")
    registrada = cliente.post(
        f"{API}/produccion",
        headers=finca["cab"],
        json={"lote_id": finca["lote"]["id"], "fecha": "2026-04-01", "detalles": [{"tipo_huevo_id": aaa["id"], "cantidad": 200}]},
    )
    assert registrada.status_code == 201, registrada.text
    antes = stock_aaa()

    # El archivo trae el mismo dia dos veces
    contenido = excel(["Fecha", "Lote", "Tipo", "Cantidad"], [["2026-04-01", "L1", "AAA", 250], ["2026-04-01", "L1", "AAA", 300]])
    analisis = subir(cliente, finca, "produccion", contenido).json()
    aplicado = cliente.post(
        f"{API}/importacion/{analisis['id']}/aplicar",
        headers=finca["cab"],
        json={"mapeo": analisis["mapeo_sugerido"], "opciones": {}},
    )
    assert aplicado.status_code == 200, aplicado.text
    assert stock_aaa() == antes + 100

    revertido = cliente.post(f"{API}/importacion/{analisis['id']}/revertir", headers=finca["cab"], json={})
    assert revertido.status_code == 200, revertido.text
    assert stock_aaa() == antes
    with SesionLocal() as db:
        dia = db.scalars(
            select(ProduccionHuevos).where(
                ProduccionHuevos.lote_id == finca["lote"]["id"],
                ProduccionHuevos.tipo_huevo_id == aaa["id"],
                ProduccionHuevos.fecha == "2026-04-01",
            )
        ).first()
        assert dia is not None and dia.cantidad == 200


def test_ninguna_ruta_bloquea_el_worker_con_trabajo_sincrono():
    """Las rutas usan la base de datos sincrona: deben ser def para correr en hilos."""
    import inspect

    from app.main import app

    asincronas = [
        ruta.path
        for ruta in app.routes
        if ruta.path.startswith("/api/v1") and inspect.iscoroutinefunction(getattr(ruta, "endpoint", None))
    ]
    assert asincronas == []


def test_leer_un_csv_no_cambia_el_dialecto_global():
    import csv

    from app.servicios.importacion import leer_archivo

    leer_archivo("solo una columna\nvalor\n".encode(), "x.csv")
    leer_archivo("a;b\n1;2\n".encode(), "y.csv")
    assert csv.excel.delimiter == ","
