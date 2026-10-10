"""Lectura de archivos de Excel o CSV y carga de sus datos al sistema.

No existe un formato fijo: el archivo se lee tal como viene, se muestran sus
columnas y el usuario dice cual corresponde a cada dato. Ese emparejamiento se
puede guardar como plantilla para la proxima vez.
"""

import csv
import io
import itertools
import re
import unicodedata
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

MAX_FILAS = 5000
MAX_TEXTO = 150  # el campo de texto mas largo que se guarda
# Topes para que un archivo hecho a proposito no tumbe el servidor
MAX_COLUMNAS = 60
MAX_FILAS_RECORRIDAS = MAX_FILAS * 4  # incluidas las vacias
MAX_DESCOMPRIMIDO = 60 * 1024 * 1024  # un .xlsx es un zip: tope a lo que ocupa abierto
MAX_PROPORCION_ZIP = 100  # mas de 100 a 1 es una bomba de descompresion
MAX_PARTES_ZIP = 2000


# --------------------------- Campos por tipo ---------------------------
# (clave, etiqueta, obligatorio, sinonimos)
CAMPOS: dict[str, list[tuple[str, str, bool, list[str]]]] = {
    "articulos": [
        ("codigo", "Codigo", True, ["codigo", "cod", "referencia", "ref", "sku"]),
        ("nombre", "Nombre", True, ["nombre", "descripcion", "articulo", "producto", "item", "insumo"]),
        ("categoria", "Categoria", False, ["categoria", "tipo", "clase", "grupo", "linea"]),
        ("unidad", "Unidad", False, ["unidad", "medida", "um", "unidad de medida", "presentacion"]),
        ("kg_por_bulto", "Kilos por bulto", False, ["kg por bulto", "kilos por bulto", "peso bulto", "kg bulto"]),
        ("stock_minimo", "Cantidad minima", False, ["minimo", "stock minimo", "existencia minima", "min"]),
        ("observaciones", "Observaciones", False, ["observaciones", "notas", "nota", "comentario"]),
    ],
    "entrada_inventario": [
        ("articulo", "Articulo", True, ["articulo", "producto", "insumo", "nombre", "descripcion", "item"]),
        ("codigo", "Codigo del articulo", False, ["codigo", "cod", "referencia", "sku"]),
        ("cantidad", "Cantidad", True, ["cantidad", "cant", "kilos", "kg", "unidades", "bultos", "existencia"]),
        ("costo_unitario", "Costo unitario", False, ["costo", "precio", "valor unitario", "costo unitario", "vr unitario"]),
        ("fecha", "Fecha", False, ["fecha", "fecha compra", "fecha factura", "dia"]),
        ("documento", "Factura o remision", False, ["factura", "documento", "remision", "numero factura", "no factura"]),
        ("proveedor", "Proveedor", False, ["proveedor", "vendedor", "distribuidor"]),
        ("lote", "Lote del producto", False, ["lote", "lote producto", "serial"]),
        ("vencimiento", "Vencimiento", False, ["vencimiento", "vence", "fecha vencimiento", "caducidad"]),
        ("unidad", "Unidad", False, ["unidad", "medida", "um"]),
        ("categoria", "Categoria", False, ["categoria", "tipo", "clase", "grupo"]),
    ],
    "proveedores": [
        ("nombre", "Nombre", True, ["nombre", "proveedor", "razon social", "empresa"]),
        ("documento", "Documento", False, ["documento", "nit", "cedula", "identificacion", "cc"]),
        ("telefono", "Telefono", False, ["telefono", "celular", "contacto", "movil"]),
        ("email", "Correo", False, ["email", "correo", "correo electronico", "mail"]),
        ("direccion", "Direccion", False, ["direccion", "ubicacion"]),
    ],
    "sanidad": [
        ("fecha", "Fecha", True, ["fecha", "dia", "fecha aplicacion"]),
        ("producto", "Producto", True, ["producto", "vacuna", "medicamento", "biologico", "nombre", "tratamiento"]),
        ("tipo", "Tipo", False, ["tipo", "clase"]),
        ("lote", "Lote de aves", False, ["lote", "lote aves", "codigo lote"]),
        ("galpon", "Galpon", False, ["galpon", "caseta", "galera"]),
        ("dosis", "Dosis", False, ["dosis", "cantidad por ave"]),
        ("via", "Via", False, ["via", "aplicacion", "metodo", "forma"]),
        ("aves_tratadas", "Aves tratadas", False, ["aves", "aves tratadas", "poblacion", "cantidad aves"]),
        ("responsable", "Responsable", False, ["responsable", "aplicado por", "encargado", "quien aplico"]),
        ("lote_producto", "Lote del producto", False, ["lote producto", "lote del producto", "serial", "batch"]),
        ("proximo_refuerzo", "Proximo refuerzo", False, ["refuerzo", "proximo refuerzo", "proxima dosis", "revacunacion"]),
        ("observaciones", "Observaciones", False, ["observaciones", "notas", "comentario"]),
    ],
    "produccion": [
        ("fecha", "Fecha", True, ["fecha", "dia"]),
        ("lote", "Lote", True, ["lote", "codigo lote", "grupo"]),
        ("tipo_huevo", "Tipo de huevo", True, ["tipo", "tipo huevo", "clasificacion", "calibre", "tamano"]),
        ("cantidad", "Cantidad", True, ["cantidad", "huevos", "unidades", "total", "recolectados"]),
    ],
}

ETIQUETAS_TIPO = {
    "articulos": "Articulos del inventario",
    "entrada_inventario": "Entrada de inventario (alimento, insumos, vacunas)",
    "proveedores": "Proveedores",
    "sanidad": "Vacunas y tratamientos aplicados",
    "produccion": "Produccion de huevos",
}


def ahora() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def normalizar(texto: str) -> str:
    """Quita tildes, espacios de mas y mayusculas para poder comparar."""
    if texto is None:
        return ""
    limpio = unicodedata.normalize("NFKD", str(texto))
    limpio = "".join(c for c in limpio if not unicodedata.combining(c))
    limpio = limpio.lower().strip()
    limpio = " ".join(limpio.replace("_", " ").replace(".", " ").split())
    return limpio


# --------------------------- Lectura del archivo ---------------------------
def _valor_plano(valor):
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return float(valor)
    if valor is None:
        return None
    if isinstance(valor, str):
        return valor.strip()
    return valor


def _nombres_de_columnas(fila: list) -> list[str]:
    """Nombres unicos para las columnas: las vacias o repetidas se numeran."""
    encabezados: list[str] = []
    for posicion, celda in enumerate(fila[:MAX_COLUMNAS], start=1):
        nombre_columna = str(celda).strip() if celda is not None and str(celda).strip() else f"Columna {posicion}"
        while nombre_columna in encabezados:
            nombre_columna = f"{nombre_columna} ({posicion})"
        encabezados.append(nombre_columna)
    return encabezados


def _revisar_zip(contenido: bytes) -> None:
    """Un .xlsx es un zip: se revisa antes de abrirlo (firma y tamano descomprimido)."""
    if not contenido.startswith(b"PK\x03\x04"):
        raise ValueError("El archivo no es un Excel .xlsx valido")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as comprimido:
            partes = comprimido.infolist()
    except zipfile.BadZipFile as error:
        raise ValueError("El archivo no es un Excel .xlsx valido") from error
    if len(partes) > MAX_PARTES_ZIP:
        raise ValueError("El archivo tiene una estructura demasiado grande")
    abierto = sum(parte.file_size for parte in partes)
    comprimido_total = sum(parte.compress_size for parte in partes) or 1
    if abierto > MAX_DESCOMPRIMIDO or abierto / comprimido_total > MAX_PROPORCION_ZIP:
        raise ValueError("El archivo es demasiado grande al abrirlo. Divide la informacion en varios archivos.")


def leer_archivo(contenido: bytes, nombre: str) -> tuple[str | None, list[str], list[dict]]:
    """Devuelve (hoja, columnas, filas). Acepta .xlsx y .csv."""
    minusculas = nombre.lower()

    if minusculas.endswith((".csv", ".txt")):
        if b"\x00" in contenido[:8192]:
            raise ValueError("El archivo no es un CSV de texto")
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            # Excel en Windows guarda los CSV en cp1252: asi no se pierden las tildes
            texto = contenido.decode("cp1252", errors="replace")
        muestra = texto[:4000]
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
        except csv.Error:
            # Una subclase propia: cambiar csv.excel afectaria a todo el proceso
            separador = ";" if muestra.count(";") > muestra.count(",") else ","
            dialecto = type("DialectoArchivo", (csv.excel,), {"delimiter": separador})

        lector = csv.reader(io.StringIO(texto), dialecto)
        # Igual que en el xlsx: se corta mientras se lee (filas vacias incluidas) y cada
        # fila se recorta a MAX_COLUMNAS. Armar todas las filas antes de recortar
        # dejaba que un CSV de 5 MB con millones de lineas agotara la memoria.
        filas_brutas = []
        for fila in itertools.islice(lector, MAX_FILAS_RECORRIDAS):
            if any(str(c).strip() for c in fila[:MAX_COLUMNAS]):
                filas_brutas.append(fila[:MAX_COLUMNAS])
                if len(filas_brutas) > MAX_FILAS:
                    break
        if not filas_brutas:
            raise ValueError("El archivo esta vacio")
        encabezados = _nombres_de_columnas(filas_brutas[0])
        datos = [
            {encabezados[i]: _valor_plano(fila[i]) if i < len(fila) else None for i in range(len(encabezados))}
            for fila in filas_brutas[1:]
        ]
        return None, encabezados, datos[:MAX_FILAS]

    if not minusculas.endswith((".xlsx", ".xlsm")):
        raise ValueError("Solo se pueden leer archivos .xlsx o .csv. Si tienes un .xls antiguo, guardalo como .xlsx")

    _revisar_zip(contenido)
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        hoja = libro.worksheets[0]
        filas_brutas = []
        for recorridas, fila in enumerate(hoja.iter_rows(values_only=True, max_col=MAX_COLUMNAS), start=1):
            # Las filas vacias tambien cuentan: una hoja con un millon de filas en blanco no se recorre entera
            if recorridas > MAX_FILAS_RECORRIDAS or len(filas_brutas) > MAX_FILAS + 20:
                break
            if fila is not None and any(celda is not None and str(celda).strip() != "" for celda in fila):
                filas_brutas.append(list(fila))
    finally:
        libro.close()

    if not filas_brutas:
        raise ValueError("La hoja esta vacia")

    # El encabezado es la primera fila con al menos dos textos
    indice_encabezado = 0
    for indice, fila in enumerate(filas_brutas[:15]):
        textos = [c for c in fila if isinstance(c, str) and c.strip()]
        if len(textos) >= 2:
            indice_encabezado = indice
            break

    encabezados = _nombres_de_columnas(filas_brutas[indice_encabezado])

    datos = []
    for fila in filas_brutas[indice_encabezado + 1 :]:
        registro = {
            encabezados[i]: _valor_plano(fila[i]) if i < len(fila) else None for i in range(len(encabezados))
        }
        if any(valor not in (None, "") for valor in registro.values()):
            datos.append(registro)

    return hoja.title, encabezados, datos[:MAX_FILAS]


def sugerir_mapeo(tipo: str, columnas: list[str]) -> dict[str, str]:
    """Empareja automaticamente las columnas del archivo con los campos del sistema."""
    sugerido: dict[str, str] = {}
    usadas: set[str] = set()
    normalizadas = {columna: normalizar(columna) for columna in columnas}

    for clave, _etiqueta, _obligatorio, sinonimos in CAMPOS.get(tipo, []):
        objetivos = [normalizar(s) for s in [clave, *sinonimos]]

        for columna, plano in normalizadas.items():
            if columna in usadas:
                continue
            if plano in objetivos:
                sugerido[clave] = columna
                usadas.add(columna)
                break

        if clave not in sugerido:
            for columna, plano in normalizadas.items():
                if columna in usadas or not plano:
                    continue
                if any(objetivo and (objetivo in plano or plano in objetivo) for objetivo in objetivos):
                    sugerido[clave] = columna
                    usadas.add(columna)
                    break

    return sugerido


# --------------------------- Conversion de valores ---------------------------
def texto(valor) -> str | None:
    if valor is None:
        return None
    limpio = str(valor).strip()
    return limpio or None


MAYOR_NUMERO = Decimal("1000000000")
# Un numero que empieza por 0 no lleva separador de miles: 0.500 es medio, no quinientos
MILES_CON_PUNTO = re.compile(r"-?[1-9]\d{0,2}(\.\d{3})+")
MILES_CON_COMA = re.compile(r"-?[1-9]\d{0,2}(,\d{3}){2,}")


def _texto_a_decimal(crudo: str) -> str:
    """Formatos de Colombia: 1.500 es mil quinientos y 1.234,56 lleva coma decimal."""
    if "," in crudo and "." in crudo:
        # El separador decimal es el que aparece de ultimo
        if crudo.rfind(",") > crudo.rfind("."):
            return crudo.replace(".", "").replace(",", ".")
        return crudo.replace(",", "")
    if "," in crudo:
        return crudo.replace(",", "") if MILES_CON_COMA.fullmatch(crudo) else crudo.replace(",", ".")
    if MILES_CON_PUNTO.fullmatch(crudo):
        return crudo.replace(".", "")
    return crudo


def numero(valor, campo: str) -> Decimal | None:
    if valor is None or str(valor).strip() == "":
        return None
    try:
        if isinstance(valor, bool):
            raise InvalidOperation
        if isinstance(valor, (int, float, Decimal)):
            # Celda numerica de Excel: ya viene como numero, no se reinterpreta
            convertido = Decimal(str(valor))
        else:
            crudo = str(valor).strip().replace("$", "").replace(" ", "")
            convertido = Decimal(_texto_a_decimal(crudo))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{campo}: '{valor}' no es un numero") from exc
    if not convertido.is_finite() or abs(convertido) >= MAYOR_NUMERO:
        raise ValueError(f"{campo}: '{valor}' no es un numero valido")
    return convertido


def fecha(valor, campo: str) -> date | None:
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    crudo = str(valor).strip()
    for formato in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(crudo, formato).date()
        except ValueError:
            continue
    raise ValueError(f"{campo}: '{valor}' no es una fecha")


def aplicar_mapeo(tipo: str, mapeo: dict[str, str], datos: dict) -> dict:
    """Pasa de {columna del archivo: valor} a {campo del sistema: valor}."""
    return {clave: datos.get(columna) for clave, columna in mapeo.items() if columna}


def validar_fila(tipo: str, valores: dict) -> dict:
    """Convierte y revisa una fila. Lanza ValueError con el problema encontrado."""
    limpio: dict = {}

    for clave, etiqueta, obligatorio, _sinonimos in CAMPOS[tipo]:
        valor = valores.get(clave)

        if clave in ("cantidad", "costo_unitario", "kg_por_bulto", "stock_minimo", "aves_tratadas"):
            convertido = numero(valor, etiqueta)
        elif clave in ("fecha", "vencimiento", "proximo_refuerzo"):
            convertido = fecha(valor, etiqueta)
        else:
            convertido = texto(valor)

        if obligatorio and convertido in (None, ""):
            raise ValueError(f"Falta {etiqueta}")
        if isinstance(convertido, Decimal) and convertido < 0:
            raise ValueError(f"{etiqueta} no puede ser negativo")
        if isinstance(convertido, str) and len(convertido) > MAX_TEXTO:
            raise ValueError(f"{etiqueta} tiene mas de {MAX_TEXTO} letras")
        limpio[clave] = convertido

    if tipo in ("entrada_inventario", "produccion"):
        if limpio.get("cantidad") is None or limpio["cantidad"] <= 0:
            raise ValueError("La cantidad debe ser mayor que cero")

    return limpio
